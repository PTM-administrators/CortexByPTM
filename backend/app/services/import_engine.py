"""
Import generico di file CSV/Excel in una tabella collegata: carichi un file,
Cortex propone come le sue colonne corrispondono a quelle della tabella
(nome uguale o simile), tu correggi se serve, e le righe valide vengono
inserite — quelle che non tornano (tipo sbagliato, colonna obbligatoria
vuota) sono segnalate riga per riga senza bloccare le altre.

A differenza della riconciliazione bancaria (bank_reconciliation.py, che
ABBINA un estratto conto a movimenti già esistenti), questo CREA righe nuove
in una tabella qualunque — è il modo con cui un'azienda porta dentro dati che
aveva altrove (es. lo storico di un Excel a 61 fogli) invece di inserirli a
mano uno a uno. Generico per qualunque settore e qualunque tabella, come il
resto della piattaforma: nessuna colonna hardcodata.
"""
import csv
import io
from typing import Any

from openpyxl import load_workbook

from app.services import data_explorer


def parse_file(nome_file: str, contenuto: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """Ritorna (intestazioni, righe) da un CSV o Excel — le righe hanno le
    chiavi = intestazioni originali del file, non ancora mappate sulla
    tabella di destinazione."""
    estensione = nome_file.rsplit(".", 1)[-1].lower() if "." in nome_file else ""
    if estensione in ("xlsx", "xlsm"):
        return _parse_excel(contenuto)
    return _parse_csv(contenuto)


def _parse_excel(contenuto: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
    ws = wb.active
    righe_grezze = ws.iter_rows(values_only=True)
    try:
        prima_riga = next(righe_grezze)
    except StopIteration:
        return [], []
    intestazioni = [str(c).strip() if c is not None else "" for c in prima_riga]

    righe = []
    for riga in righe_grezze:
        if all(v is None for v in riga):
            continue  # riga vuota, capita spesso a fine foglio Excel
        righe.append({intestazioni[i]: riga[i] for i in range(len(intestazioni)) if i < len(riga)})
    return intestazioni, righe


def _parse_csv(contenuto: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    testo = contenuto.decode("utf-8-sig")  # utf-8-sig: tollera il BOM che Excel aggiunge ai CSV
    campione = testo[:2000]
    try:
        dialetto = csv.Sniffer().sniff(campione, delimiters=",;\t")
    except csv.Error:
        dialetto = csv.excel  # separatore non riconosciuto (es. una sola colonna): virgola di default

    lettore = csv.DictReader(io.StringIO(testo), dialect=dialetto)
    intestazioni = [(h or "").strip() for h in (lettore.fieldnames or [])]
    righe = [{(k or "").strip(): v for k, v in riga.items()} for riga in lettore]
    return intestazioni, righe


def _normalizza(testo: str) -> str:
    return "".join(c for c in testo.lower().strip() if c.isalnum())


def suggerisci_mapping(intestazioni_file: list[str], colonne_tabella: list[str]) -> dict[str, str]:
    """Colonna della tabella -> intestazione del file, indovinata per
    corrispondenza esatta del nome (normalizzato, es. "Importo €" e
    "importo" combaciano) — stesso principio di "indovina" già usato per le
    Viste e le Segnalazioni, qui applicato al caso più semplice di colonna
    contro colonna invece che colonna contro ruolo. Le colonne senza una
    corrispondenza restano fuori dal mapping: l'utente le assegna a mano se
    vuole, non si inventa un abbinamento incerto."""
    normalizzate = {_normalizza(h): h for h in intestazioni_file}
    mapping = {}
    for colonna in colonne_tabella:
        corrispondenza = normalizzate.get(_normalizza(colonna))
        if corrispondenza:
            mapping[colonna] = corrispondenza
    return mapping


def prepara_righe(righe_file: list[dict[str, Any]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    """Applica il mapping colonna-tabella -> intestazione-file: da qui in poi
    le righe hanno le chiavi della TABELLA, come si aspetta data_explorer."""
    return [{colonna: riga.get(intestazione) for colonna, intestazione in mapping.items()} for riga in righe_file]


def importa(connection_string: str, table_name: str, righe: list[dict[str, Any]], esegui: bool) -> dict[str, Any]:
    """Valida (e se esegui=True inserisce davvero) ogni riga indipendentemente
    dalle altre: una riga con un valore sbagliato non blocca le righe buone,
    viene solo segnalata con il proprio numero e il motivo."""
    valide = 0
    inserite = 0
    errori: list[dict[str, Any]] = []

    for indice, riga in enumerate(righe, start=1):
        payload = {k: v for k, v in riga.items() if v not in (None, "")}
        if not payload:
            continue  # riga vuota dopo il mapping (es. una riga in bianco nel file): ignorata, non un errore
        try:
            if esegui:
                data_explorer.insert_row(connection_string, table_name, payload)
                inserite += 1
            else:
                data_explorer.validate_row(connection_string, table_name, payload)
            valide += 1
        except Exception as exc:
            errori.append({"riga": indice, "errore": str(exc)})

    return {
        "totale": len(righe),
        "valide": valide,
        "errori": errori,
        "inserite": inserite if esegui else None,
    }
